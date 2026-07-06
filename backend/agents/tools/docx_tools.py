from __future__ import annotations

import base64
import copy
import io
import json
import os
import re
import shutil
import subprocess
import zipfile
from typing import Any, Optional
from xml.etree import ElementTree as ET

from backend.agents.tools.workspace_tools import get_safe_workspace_path


def materialize_original_resume_file(
    user_id: Any,
    task_id: str,
    original_resume_id: Optional[str],
    *,
    db: Any = None,
) -> Optional[dict[str, str]]:
    """
    将用户上传的原始简历文件写入任务工作区，供后续高保真替换/导出使用。
    目前只有 DOCX 能进行可编辑高保真替换；PDF 原件也会保留，避免误判为可降级重建。
    """
    if not original_resume_id:
        return None

    if db is None:
        from database import Database

        db = Database()

    resume_data = db.get_user_resume(user_id, original_resume_id)

    if not resume_data or not isinstance(resume_data.get("file"), dict):
        return None

    file_info = resume_data["file"]
    b64_data = str(file_info.get("b64_content") or "")
    file_name = str(file_info.get("name") or "")
    suffix = os.path.splitext(file_name)[1].lower()
    if suffix not in {".docx", ".pdf"} or not b64_data:
        return None

    try:
        encoded = b64_data.split(",", 1)[1] if "," in b64_data else b64_data
        file_bytes = base64.b64decode(encoded)
    except Exception as exc:
        print(f"[ORIGINAL_RESUME] Failed to decode original file: {exc}")
        return None

    source_format = suffix.lstrip(".")
    target_name = f"original_resume.{source_format}"
    target_path = get_safe_workspace_path(user_id, task_id, target_name)
    with open(target_path, "wb") as f:
        f.write(file_bytes)

    return {"source_format": source_format, "path": target_path}


def extract_docx_style_profile(user_id: Any, task_id: str, original_resume_id: Optional[str]) -> None:
    """
    从用户上传的原始简历中提取格式信息并保存到工作区中的 style_profile.json。
    """
    from database import Database

    profile = {
        "body_font": "SimSun",
        "body_size": 11.0,
        "line_spacing": 1.15,
        "space_after": 6.0,
        "space_before": 0.0,
        "margins": {
            "top": 2.0,
            "bottom": 2.0,
            "left": 2.5,
            "right": 2.5,
        },
        "heading_font": "SimSun",
        "heading_size": 14.0,
    }

    if original_resume_id:
        db = Database()
        resume_data = db.get_user_resume(user_id, original_resume_id)
        if resume_data and "file" in resume_data:
            b64_data = resume_data["file"].get("b64_content", "")
            file_name = resume_data["file"].get("name", "")
            suffix = os.path.splitext(file_name)[1].lower()

            if b64_data:
                try:
                    file_bytes = base64.b64decode(b64_data)
                    if suffix == ".docx":
                        orig_docx_path = get_safe_workspace_path(user_id, task_id, "original_resume.docx")
                        with open(orig_docx_path, "wb") as f:
                            f.write(file_bytes)
                        import docx

                        doc = docx.Document(io.BytesIO(file_bytes))
                        body_fonts: dict[str, int] = {}
                        body_sizes: dict[float, int] = {}
                        heading_fonts: dict[str, int] = {}
                        heading_sizes: dict[float, int] = {}

                        for paragraph in doc.paragraphs:
                            is_heading = paragraph.style.name.startswith("Heading") or paragraph.style.name.startswith("Title")
                            p_format = paragraph.paragraph_format
                            if p_format.line_spacing is not None and isinstance(p_format.line_spacing, (int, float)):
                                profile["line_spacing"] = round(p_format.line_spacing, 2)
                            if p_format.space_after is not None:
                                profile["space_after"] = round(p_format.space_after.pt, 1)
                            if p_format.space_before is not None:
                                profile["space_before"] = round(p_format.space_before.pt, 1)

                            for run in paragraph.runs:
                                if run.font.name:
                                    target = heading_fonts if is_heading else body_fonts
                                    target[run.font.name] = target.get(run.font.name, 0) + 1
                                if run.font.size:
                                    size = run.font.size.pt
                                    target = heading_sizes if is_heading else body_sizes
                                    target[size] = target.get(size, 0) + 1

                        if body_fonts:
                            profile["body_font"] = max(body_fonts, key=body_fonts.get)
                        if body_sizes:
                            profile["body_size"] = max(body_sizes, key=body_sizes.get)
                        if heading_fonts:
                            profile["heading_font"] = max(heading_fonts, key=heading_fonts.get)
                        if heading_sizes:
                            profile["heading_size"] = max(heading_sizes, key=heading_sizes.get)

                        if doc.sections:
                            section = doc.sections[0]
                            profile["margins"] = {
                                "top": round(section.top_margin.cm, 2) if section.top_margin else 2.0,
                                "bottom": round(section.bottom_margin.cm, 2) if section.bottom_margin else 2.0,
                                "left": round(section.left_margin.cm, 2) if section.left_margin else 2.5,
                                "right": round(section.right_margin.cm, 2) if section.right_margin else 2.5,
                            }
                    elif suffix == ".pdf":
                        import pdfplumber

                        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                            body_fonts: dict[str, int] = {}
                            body_sizes: dict[float, int] = {}
                            for page in pdf.pages:
                                if not getattr(page, "chars", None):
                                    continue
                                for obj in page.chars:
                                    font = obj.get("fontname", "")
                                    size = obj.get("size", 11)
                                    if font:
                                        clean_font = font.split("+")[-1].split("-")[0]
                                        body_fonts[clean_font] = body_fonts.get(clean_font, 0) + 1
                                    if size:
                                        rounded_size = round(size * 2) / 2
                                        body_sizes[rounded_size] = body_sizes.get(rounded_size, 0) + 1

                            if body_fonts:
                                profile["body_font"] = max(body_fonts, key=body_fonts.get)
                            if body_sizes:
                                sorted_sizes = sorted(body_sizes.items(), key=lambda item: item[1], reverse=True)
                                profile["body_size"] = sorted_sizes[0][0]
                                headings = [
                                    size
                                    for size, count in sorted_sizes
                                    if size > profile["body_size"] and count > sorted_sizes[0][1] * 0.05
                                ]
                                if headings:
                                    profile["heading_size"] = max(headings)
                                    profile["heading_font"] = profile["body_font"]
                except Exception as exc:
                    print(f"[STYLE_EXTRACT_ERROR] Failed to extract style: {exc}")

    try:
        profile_path = get_safe_workspace_path(user_id, task_id, "style_profile.json")
        with open(profile_path, "w", encoding="utf-8") as f:
            json.dump(profile, f, ensure_ascii=False, indent=2)
    except Exception as exc:
        print(f"[STYLE_SAVE_ERROR] Failed to save style profile: {exc}")


def generate_docx_from_markdown(user_id: Any, task_id: str, markdown_content: str, docx_path: str) -> None:
    """
    根据 style_profile.json 生成样式保真的 DOCX 简历文件。
    """
    import docx
    from docx.shared import Cm, Pt

    profile = {
        "body_font": "SimSun",
        "body_size": 11,
        "line_spacing": 1.15,
        "space_after": 6.0,
        "space_before": 0.0,
        "margins": {
            "top": 2.0,
            "bottom": 2.0,
            "left": 2.5,
            "right": 2.5,
        },
        "heading_font": "SimSun",
        "heading_size": 14,
    }

    try:
        profile_path = get_safe_workspace_path(user_id, task_id, "style_profile.json")
        if os.path.exists(profile_path):
            with open(profile_path, "r", encoding="utf-8") as f:
                loaded = json.load(f)
                if isinstance(loaded, dict):
                    profile.update(loaded)
    except Exception as exc:
        print(f"[DOCX_GEN_WARN] Failed to load style profile: {exc}")

    doc = docx.Document()
    if doc.sections:
        section = doc.sections[0]
        margins = profile.get("margins", {})
        section.top_margin = Cm(margins.get("top", 2.0))
        section.bottom_margin = Cm(margins.get("bottom", 2.0))
        section.left_margin = Cm(margins.get("left", 2.5))
        section.right_margin = Cm(margins.get("right", 2.5))

    blocks = re.split(r"\n\n+", markdown_content)
    for block in blocks:
        block = block.strip()
        if not block:
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", block)
        if heading_match:
            level = len(heading_match.group(1))
            title_text = heading_match.group(2)
            paragraph = doc.add_paragraph()
            p_format = paragraph.paragraph_format
            p_format.space_before = Pt(12)
            p_format.space_after = Pt(6)
            p_format.keep_with_next = True
            run = paragraph.add_run(title_text)
            run.bold = True
            run.font.name = profile.get("heading_font", "SimSun")
            run._r.get_or_add_rPr().get_or_add_rFonts().set(
                docx.oxml.ns.qn("w:eastAsia"),
                profile.get("heading_font", "SimSun"),
            )
            heading_size = profile.get("heading_size", 14)
            if level > 1:
                heading_size = max(heading_size - 2, 12)
            run.font.size = Pt(heading_size)
            continue

        list_match = re.match(r"^[\-\*\+]\s+(.*)$", block)
        if list_match:
            for line in block.splitlines():
                line_match = re.match(r"^[\-\*\+]\s+(.*)$", line.strip())
                paragraph = doc.add_paragraph(style="List Bullet" if line_match else None)
                p_format = paragraph.paragraph_format
                if not line_match:
                    p_format.left_indent = Cm(0.74)
                p_format.line_spacing = profile.get("line_spacing", 1.15)
                p_format.space_after = Pt(profile.get("space_after", 6.0))
                p_format.space_before = Pt(profile.get("space_before", 0.0))
                add_markdown_runs(paragraph, line_match.group(1) if line_match else line.strip(), profile)
            continue

        paragraph = doc.add_paragraph()
        p_format = paragraph.paragraph_format
        p_format.line_spacing = profile.get("line_spacing", 1.15)
        p_format.space_after = Pt(profile.get("space_after", 6.0))
        p_format.space_before = Pt(profile.get("space_before", 0.0))
        add_markdown_runs(paragraph, block, profile)

    doc.save(docx_path)


def add_markdown_runs(paragraph: Any, text: str, profile: dict[str, Any]) -> None:
    import docx
    from docx.shared import Pt

    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run_text = part[2:-2]
            is_bold = True
        else:
            run_text = part
            is_bold = False
        if not run_text:
            continue
        run = paragraph.add_run(run_text)
        run.bold = is_bold
        run.font.name = profile.get("body_font", "SimSun")
        run._r.get_or_add_rPr().get_or_add_rFonts().set(
            docx.oxml.ns.qn("w:eastAsia"),
            profile.get("body_font", "SimSun"),
        )
        run.font.size = Pt(profile.get("body_size", 11))


def convert_docx_to_pdf(user_id: Any, task_id: str, docx_path: str, pdf_dir: str) -> bool:
    """
    将 DOCX 转换为 PDF。优先使用 LibreOffice，回退到 docx2pdf。
    """
    expected_pdf_path = os.path.join(
        pdf_dir,
        os.path.splitext(os.path.basename(docx_path))[0] + ".pdf",
    )
    legacy_pdf_path = os.path.join(pdf_dir, "optimized_resume.pdf")
    libreoffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
    if libreoffice_bin:
        try:
            print(f"[PDF_CONVERT] Found LibreOffice at {libreoffice_bin}. Starting headless conversion...")
            result = subprocess.run(
                [libreoffice_bin, "--headless", "--convert-to", "pdf", "--outdir", pdf_dir, docx_path],
                capture_output=True,
                text=True,
                timeout=30,
            )
            if result.returncode == 0 and (os.path.exists(expected_pdf_path) or os.path.exists(legacy_pdf_path)):
                print("[PDF_CONVERT] LibreOffice conversion succeeded.")
                return True
            print(f"[PDF_CONVERT] LibreOffice return code: {result.returncode}, stderr: {result.stderr}")
        except Exception as exc:
            print(f"[PDF_CONVERT] LibreOffice failed with exception: {exc}")

    try:
        from docx2pdf import convert

        print("[PDF_CONVERT] LibreOffice not found or failed. Falling back to docx2pdf...")
        pdf_path = legacy_pdf_path
        try:
            import pythoncom

            pythoncom.CoInitialize()
        except Exception:
            pass
        convert(docx_path, pdf_path)
        if os.path.exists(pdf_path):
            print("[PDF_CONVERT] docx2pdf conversion succeeded.")
            return True
    except Exception as exc:
        print(f"[PDF_CONVERT] docx2pdf fallback failed: {exc}")

    return False


_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_VML_NS = "urn:schemas-microsoft-com:vml"

_DOCX_TEMPLATE_GUARD_FIELDS = (
    "layout_xml_parts",
    "paragraphs",
    "tables",
    "table_cells",
    "drawings",
    "picts",
    "vml_shapes",
    "vml_textboxes",
    "textbox_contents",
    "header_parts",
    "footer_parts",
    "media_files",
    "image_relationships",
)


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _xml_namespace(tag: str) -> str:
    if tag.startswith("{") and "}" in tag:
        return tag[1:].split("}", 1)[0]
    return ""


def _is_layout_xml_part(name: str) -> bool:
    if name == "word/document.xml":
        return True
    base_name = os.path.basename(name)
    return (
        name.startswith("word/")
        and name.endswith(".xml")
        and (
            base_name.startswith("header")
            or base_name.startswith("footer")
            or base_name in {"footnotes.xml", "endnotes.xml", "comments.xml"}
        )
    )


def collect_docx_template_fingerprint(docx_path: str) -> dict[str, int]:
    fingerprint = {
        "package_parts": 0,
        "layout_xml_parts": 0,
        "relationship_parts": 0,
        "paragraphs": 0,
        "tables": 0,
        "table_cells": 0,
        "drawings": 0,
        "picts": 0,
        "vml_shapes": 0,
        "vml_textboxes": 0,
        "textbox_contents": 0,
        "header_parts": 0,
        "footer_parts": 0,
        "media_files": 0,
        "image_relationships": 0,
    }

    with zipfile.ZipFile(docx_path, "r") as archive:
        names = [name for name in archive.namelist() if not name.endswith("/")]
        fingerprint["package_parts"] = len(names)
        fingerprint["layout_xml_parts"] = sum(1 for name in names if _is_layout_xml_part(name))
        fingerprint["relationship_parts"] = sum(1 for name in names if name.endswith(".rels"))
        fingerprint["header_parts"] = sum(1 for name in names if name.startswith("word/header") and name.endswith(".xml"))
        fingerprint["footer_parts"] = sum(1 for name in names if name.startswith("word/footer") and name.endswith(".xml"))
        fingerprint["media_files"] = sum(1 for name in names if name.startswith("word/media/"))

        for name in names:
            if not _is_layout_xml_part(name):
                continue
            root = ET.fromstring(archive.read(name))
            for element in root.iter():
                namespace = _xml_namespace(element.tag)
                local_name = _xml_local_name(element.tag)
                if namespace == _WORD_NS:
                    if local_name == "p":
                        fingerprint["paragraphs"] += 1
                    elif local_name == "tbl":
                        fingerprint["tables"] += 1
                    elif local_name == "tc":
                        fingerprint["table_cells"] += 1
                    elif local_name == "drawing":
                        fingerprint["drawings"] += 1
                    elif local_name == "pict":
                        fingerprint["picts"] += 1
                    elif local_name == "txbxContent":
                        fingerprint["textbox_contents"] += 1
                elif namespace == _VML_NS:
                    if local_name == "shape":
                        fingerprint["vml_shapes"] += 1
                    elif local_name == "textbox":
                        fingerprint["vml_textboxes"] += 1

        for name in names:
            if not name.endswith(".rels"):
                continue
            root = ET.fromstring(archive.read(name))
            for element in root.iter():
                if _xml_local_name(element.tag) != "Relationship":
                    continue
                rel_type = str(element.attrib.get("Type") or "")
                if rel_type.endswith("/image") or "/image" in rel_type:
                    fingerprint["image_relationships"] += 1

    return fingerprint


def validate_docx_template_preservation(
    before: dict[str, int],
    after: dict[str, int],
) -> dict[str, Any]:
    issues = []
    for field in _DOCX_TEMPLATE_GUARD_FIELDS:
        before_value = int(before.get(field, 0) or 0)
        after_value = int(after.get(field, 0) or 0)
        if after_value < before_value:
            issues.append({
                "field": field,
                "before": before_value,
                "after": after_value,
                "message": f"{field} decreased from {before_value} to {after_value}.",
            })

    return {
        "ok": len(issues) == 0,
        "issues": issues,
        "before": before,
        "after": after,
    }


def _write_docx_template_guard_report(user_id: Any, task_id: str, report: dict[str, Any]) -> None:
    report_path = get_safe_workspace_path(user_id, task_id, "docx_template_guard.json")
    with open(report_path, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)


def _remove_unverified_docx(docx_path: str) -> None:
    try:
        if os.path.exists(docx_path):
            os.remove(docx_path)
    except Exception as exc:
        print(f"[DOCX_TEMPLATE_GUARD] Failed to remove unverified DOCX: {exc}")


def _fail_docx_template_guard(
    user_id: Any,
    task_id: str,
    opt_docx_path: str,
    reason: str,
    *,
    issues: list[dict[str, Any]] | None = None,
    before: dict[str, int] | None = None,
    after: dict[str, int] | None = None,
) -> bool:
    report = {
        "ok": False,
        "reason": reason,
        "issues": issues or [],
        "before": before or {},
        "after": after or {},
    }
    _write_docx_template_guard_report(user_id, task_id, report)
    _remove_unverified_docx(opt_docx_path)
    print(f"[DOCX_TEMPLATE_GUARD] {reason}")
    return False


def update_docx_resume_from_log(user_id: Any, task_id: str) -> bool:
    """
    用 modification_log.json 中的修改记录对 original_resume.docx 原地替换，
    生成高保真的 optimized_resume.docx，并对表格/文本框做防溢出处理。
    """
    import docx

    orig_docx_path = get_safe_workspace_path(user_id, task_id, "original_resume.docx")
    opt_docx_path = get_safe_workspace_path(user_id, task_id, "optimized_resume.docx")
    mod_log_path = get_safe_workspace_path(user_id, task_id, "modification_log.json")

    if not os.path.exists(orig_docx_path):
        print(f"[DOCX_REPLACE] original_resume.docx not found at {orig_docx_path}. Skipping high-fidelity replacement.")
        _remove_unverified_docx(opt_docx_path)
        return False
    if not os.path.exists(mod_log_path):
        print(f"[DOCX_REPLACE] modification_log.json not found at {mod_log_path}. Skipping.")
        _remove_unverified_docx(opt_docx_path)
        return False

    source_fingerprint: dict[str, int] | None = None
    try:
        source_fingerprint = collect_docx_template_fingerprint(orig_docx_path)
        shutil.copyfile(orig_docx_path, opt_docx_path)
        with open(mod_log_path, "r", encoding="utf-8") as f:
            modification_log = json.load(f)
        if not modification_log:
            print("[DOCX_REPLACE] modification_log.json is empty. No replacements needed.")
            copied_fingerprint = collect_docx_template_fingerprint(opt_docx_path)
            guard_report = validate_docx_template_preservation(source_fingerprint, copied_fingerprint)
            if not guard_report["ok"]:
                guard_report["reason"] = "DOCX copy failed template preservation checks."
                _write_docx_template_guard_report(user_id, task_id, guard_report)
                _remove_unverified_docx(opt_docx_path)
                return False
            return True

        doc = docx.Document(opt_docx_path)
        w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        v_ns = "urn:schemas-microsoft-com:vml"
        w_txbx_content_tag = f"{{{w_ns}}}txbxContent"
        w_paragraph_tag = f"{{{w_ns}}}p"
        v_shape_tag = f"{{{v_ns}}}shape"

        def get_all_paragraphs_with_textboxes(parent_doc: Any) -> list[Any]:
            paragraphs: list[Any] = []
            if hasattr(parent_doc, "paragraphs"):
                paragraphs.extend(parent_doc.paragraphs)
            if hasattr(parent_doc, "sections"):
                for section in parent_doc.sections:
                    for part in (section.header, section.footer):
                        paragraphs.extend(get_all_paragraphs_with_textboxes(part))
            if hasattr(parent_doc, "tables"):
                for table in parent_doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            paragraphs.extend(get_all_paragraphs_with_textboxes(cell))
            element = getattr(parent_doc, "element", None)
            if element is not None:
                try:
                    for txbx in element.iter(w_txbx_content_tag):
                        for p_el in txbx.iter(w_paragraph_tag):
                            paragraphs.append(docx.text.paragraph.Paragraph(p_el, doc))
                except Exception as exc:
                    print(f"[DOCX_REPLACE] Error searching textboxes: {exc}")
            return paragraphs

        def clean_text(text: str) -> str:
            return "".join(text.split()) if text else ""

        def unique_paragraphs(paragraphs: list[Any]) -> list[Any]:
            unique: list[Any] = []
            seen: set[int] = set()
            for paragraph in paragraphs:
                marker = id(getattr(paragraph, "_p", paragraph))
                if marker in seen:
                    continue
                seen.add(marker)
                unique.append(paragraph)
            return unique

        def contains_textbox_content(paragraph: Any) -> bool:
            p_el = getattr(paragraph, "_p", None)
            if p_el is None:
                return False
            try:
                return any(True for _ in p_el.iter(w_txbx_content_tag))
            except Exception:
                return False

        def capture_run_style(paragraph: Any) -> dict[str, Any]:
            if not paragraph.runs:
                return {}
            base_run = next((run for run in paragraph.runs if run.text.strip()), paragraph.runs[0])
            color = None
            try:
                color = base_run.font.color.rgb if base_run.font.color and base_run.font.color.rgb else None
            except Exception:
                color = None
            return {
                "font_name": base_run.font.name,
                "font_size": base_run.font.size,
                "font_color": color,
                "bold": base_run.font.bold,
                "italic": base_run.font.italic,
                "underline": base_run.font.underline,
            }

        def apply_run_style(run: Any, style: dict[str, Any]) -> None:
            font_name = style.get("font_name")
            if font_name:
                run.font.name = font_name
                try:
                    run._r.get_or_add_rPr().get_or_add_rFonts().set(docx.oxml.ns.qn("w:eastAsia"), font_name)
                except Exception:
                    pass
            if style.get("font_size"):
                run.font.size = style["font_size"]
            if style.get("font_color"):
                run.font.color.rgb = style["font_color"]
            for attr in ("bold", "italic", "underline"):
                if style.get(attr) is not None:
                    setattr(run.font, attr, style[attr])

        def set_paragraph_text_preserving_style(paragraph: Any, text: str) -> None:
            alignment = paragraph.alignment
            style = capture_run_style(paragraph)
            if hasattr(paragraph, "clear"):
                paragraph.clear()
            else:
                paragraph.text = ""
            if text:
                lines = text.splitlines() or [text]
                run = paragraph.add_run(lines[0])
                apply_run_style(run, style)
                for line in lines[1:]:
                    run.add_break()
                    run.add_text(line)
            if alignment is not None:
                paragraph.alignment = alignment

        def split_replacement_lines(text: str, target_count: int) -> list[str]:
            lines = [line.strip() for line in str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
            lines = [line for line in lines if line]
            if not lines:
                return ["" for _ in range(max(target_count, 1))]
            if len(lines) <= target_count:
                return lines + ["" for _ in range(target_count - len(lines))]
            return lines

        def clone_paragraph_after(paragraph: Any) -> Any:
            parent = getattr(paragraph, "_parent", doc)
            new_p = copy.deepcopy(paragraph._p)
            paragraph._p.addnext(new_p)
            cloned = docx.text.paragraph.Paragraph(new_p, parent)
            set_paragraph_text_preserving_style(cloned, "")
            return cloned

        def replace_paragraph_group(paragraphs: list[Any], new_text: str) -> None:
            if not paragraphs:
                return
            replacement_lines = split_replacement_lines(new_text, len(paragraphs))
            while len(paragraphs) < len(replacement_lines):
                paragraphs.append(clone_paragraph_after(paragraphs[-1]))
            for paragraph, line in zip(paragraphs, replacement_lines):
                set_paragraph_text_preserving_style(paragraph, line)

        def is_single_paragraph_match(original_clean: str, paragraph_clean: str) -> bool:
            if not original_clean or not paragraph_clean:
                return False
            if original_clean == paragraph_clean:
                return True
            if original_clean in paragraph_clean:
                return True
            if paragraph_clean in original_clean:
                return len(paragraph_clean) >= max(30, int(len(original_clean) * 0.8))
            return False

        def is_window_match(original_clean: str, combined_clean: str) -> bool:
            if not original_clean or not combined_clean:
                return False
            if original_clean == combined_clean or original_clean in combined_clean:
                return True
            if combined_clean in original_clean:
                return False
            coverage = min(len(original_clean), len(combined_clean)) / max(len(original_clean), len(combined_clean))
            if coverage < 0.7:
                return False
            from difflib import SequenceMatcher

            return SequenceMatcher(None, original_clean, combined_clean).ratio() >= 0.92

        all_paragraphs = unique_paragraphs(get_all_paragraphs_with_textboxes(doc))
        paragraph_entries = [
            {"paragraph": paragraph, "clean": clean_text(paragraph.text), "used": False}
            for paragraph in all_paragraphs
            if clean_text(paragraph.text) and not contains_textbox_content(paragraph)
        ]

        matched_count = 0
        for item in modification_log:
            if not isinstance(item, dict):
                continue
            original_clean = clean_text(item.get("original", ""))
            new_text = str(item.get("new", ""))
            if not original_clean or not new_text:
                continue

            matched_entries: Optional[list[dict[str, Any]]] = None
            for entry in paragraph_entries:
                if entry["used"]:
                    continue
                if is_single_paragraph_match(original_clean, entry["clean"]):
                    matched_entries = [entry]
                    break

            if matched_entries is None:
                max_window_size = 160
                for start_idx, entry in enumerate(paragraph_entries):
                    if entry["used"]:
                        continue
                    combined = ""
                    window: list[dict[str, Any]] = []
                    for entry in paragraph_entries[start_idx:]:
                        if entry["used"]:
                            break
                        combined += entry["clean"]
                        window.append(entry)
                        if is_window_match(original_clean, combined):
                            matched_entries = window
                            break
                        if len(window) >= max_window_size:
                            break
                        if len(combined) > len(original_clean) * 1.6 and original_clean not in combined:
                            break
                    if matched_entries is not None:
                        break

            if matched_entries is None:
                print(f"[DOCX_REPLACE] No DOCX paragraph match for section {item.get('section_name') or item.get('section_index')}.")
                continue

            replace_paragraph_group([entry["paragraph"] for entry in matched_entries], new_text)
            for entry in matched_entries:
                entry["used"] = True
                entry["clean"] = clean_text(entry["paragraph"].text)
            matched_count += 1

        print(f"[DOCX_REPLACE] Replaced {matched_count} modification items in docx.")
        if matched_count == 0 and modification_log:
            after_fingerprint = collect_docx_template_fingerprint(opt_docx_path) if os.path.exists(opt_docx_path) else {}
            return _fail_docx_template_guard(
                user_id,
                task_id,
                opt_docx_path,
                "No DOCX paragraph matched modification_log; high-fidelity DOCX was not generated.",
                issues=[{
                    "field": "paragraph_matches",
                    "before": len(modification_log),
                    "after": matched_count,
                    "message": "No modification item matched the original DOCX text.",
                }],
                before=source_fingerprint,
                after=after_fingerprint,
            )

        try:
            for shape in doc.element.iter(v_shape_tag):
                if any(True for _ in shape.iter(w_txbx_content_tag)):
                    style = shape.get("style")
                    if style:
                        parts = []
                        has_fit = False
                        for part in style.split(";"):
                            part = part.strip()
                            if not part:
                                continue
                            if part.startswith("height:"):
                                parts.append("height:auto")
                            elif part.startswith("mso-fit-shape-to-text:"):
                                parts.append("mso-fit-shape-to-text:t")
                                has_fit = True
                            else:
                                parts.append(part)
                        if not has_fit:
                            parts.append("mso-fit-shape-to-text:t")
                        shape.set("style", ";".join(parts))
            print("[DOCX_REPLACE] Adjusted height attributes of VML textboxes.")
        except Exception as exc:
            print(f"[DOCX_REPLACE] Error processing VML textbox heights: {exc}")

        try:
            for table in doc.tables:
                for row in table.rows:
                    tr_pr = row._tr.get_or_add_trPr()
                    for child in list(tr_pr):
                        if child.tag == docx.oxml.ns.qn("w:cantSplit"):
                            tr_pr.remove(child)
                        elif (
                            child.tag == docx.oxml.ns.qn("w:trHeight")
                            and child.get(docx.oxml.ns.qn("w:hRule")) == "exact"
                        ):
                            child.set(docx.oxml.ns.qn("w:hRule"), "atLeast")
            print("[DOCX_REPLACE] Adjusted cantSplit and height rules for tables.")
        except Exception as exc:
            print(f"[DOCX_REPLACE] Error processing table rows: {exc}")

        doc.save(opt_docx_path)
        saved_fingerprint = collect_docx_template_fingerprint(opt_docx_path)
        guard_report = validate_docx_template_preservation(source_fingerprint, saved_fingerprint)
        if not guard_report["ok"]:
            guard_report["reason"] = "DOCX template structure changed after save; high-fidelity output rejected."
            _write_docx_template_guard_report(user_id, task_id, guard_report)
            _remove_unverified_docx(opt_docx_path)
            print("[DOCX_TEMPLATE_GUARD] DOCX template structure changed after save.")
            return False
        print(f"[DOCX_REPLACE] High-fidelity docx saved to {opt_docx_path}")
        return True
    except Exception as exc:
        import traceback

        print(f"[DOCX_REPLACE_ERROR] {exc}\n{traceback.format_exc()}")
        return _fail_docx_template_guard(
            user_id,
            task_id,
            opt_docx_path,
            f"DOCX replacement failed: {exc}",
            before=source_fingerprint,
        )
