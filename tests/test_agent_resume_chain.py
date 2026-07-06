import os
import json
import pytest
from pathlib import Path
from backend.agent_resume import (
    tool_extract_resume_sections,
    tool_replace_resume_section,
    tool_generate_modification_diff,
    get_safe_workspace_path
)
from backend.agents.tools.docx_tools import (
    collect_docx_template_fingerprint,
    generate_docx_from_markdown,
    materialize_original_resume_file,
    update_docx_resume_from_log,
    validate_docx_template_preservation,
)
from backend.main import create_app
from config import Config
from database import Database
from service import CareerPathAIService
from fastapi.testclient import TestClient

# Mock sample resume
SAMPLE_RESUME = """张三
联系电话：13800000000
邮箱：zhangsan@example.com

教育经历
北京大学 - 计算机科学与技术 - 本科 - 2020.09-2024.07

技能
Python, FastAPI, React, Docker

项目经历
项目A：个人求职看板 InternPath
负责实现 AI Agent 简历优化链，重构了大模型交互逻辑。
"""

def test_tool_extract_resume_sections():
    sections_json = tool_extract_resume_sections(SAMPLE_RESUME)
    sections = json.loads(sections_json)

    assert len(sections) > 0
    # First section should be '其他' containing contact info
    assert sections[0]["section_name"] == "其他"
    assert "张三" in sections[0]["content"]
    assert sections[0]["index"] == 0

    # Verify '教育经历' is parsed
    edu_section = next((s for s in sections if s["section_name"] == "教育经历"), None)
    assert edu_section is not None
    assert "北京大学" in edu_section["content"]

    # Verify '项目经历' is parsed
    proj_section = next((s for s in sections if s["section_name"] == "项目经历"), None)
    assert proj_section is not None
    assert "InternPath" in proj_section["content"]


def test_tool_replace_and_diff(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))
    user_id = 42
    task_id = "test_task"

    # Setup workspace by writing resume_sections.json
    workspace_dir = Path(get_safe_workspace_path(user_id, task_id, "temp.txt")).parent
    sections_json = tool_extract_resume_sections(SAMPLE_RESUME)

    with open(workspace_dir / "resume_sections.json", "w", encoding="utf-8") as f:
        f.write(sections_json)

    # Perform replacement
    res = tool_replace_resume_section(
        user_id=user_id,
        task_id=task_id,
        section_index=2, # index 2 is '技能' (0 is 其他, 1 is 教育经历, 2 is 技能)
        new_content="技能\nPython, Go, Kubernetes, Terraform",
        reason="突出容器化和云原生运维能力"
    )

    assert "成功" in res

    # Read assembled resume
    with open(workspace_dir / "assembled_resume.txt", "r", encoding="utf-8") as f:
        assembled = f.read()
    assert "Go, Kubernetes" in assembled
    assert "北京大学" in assembled # unchanged parts must remain

    # Generate diff
    diff = tool_generate_modification_diff(user_id, task_id)
    assert "突出容器化和云原生运维能力" in diff
    assert "Python, FastAPI" in diff # original text
    assert "Go, Kubernetes" in diff # new text


def test_workspace_path_rejects_traversal_and_absolute_paths(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))
    user_id = 42
    task_id = "safe-path-task"

    safe_path = Path(get_safe_workspace_path(user_id, task_id, "resume_sections.json"))
    assert safe_path.parent.name == f"task_{task_id}"

    with pytest.raises(PermissionError):
        get_safe_workspace_path(user_id, task_id, "../escape.txt")

    prefix_sibling = f"../task_{task_id}_evil/escape.txt"
    with pytest.raises(PermissionError):
        get_safe_workspace_path(user_id, task_id, prefix_sibling)

    with pytest.raises(PermissionError):
        get_safe_workspace_path(user_id, task_id, str(tmp_path / "absolute.txt"))


def test_docx_tool_generates_docx_from_markdown(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))
    user_id = 42
    task_id = "docx-tool-task"
    docx_path = get_safe_workspace_path(user_id, task_id, "optimized_resume.docx")

    generate_docx_from_markdown(
        user_id=user_id,
        task_id=task_id,
        markdown_content="# 简历\n\n- 负责开发 InternPath。",
        docx_path=docx_path,
    )

    assert Path(docx_path).exists()
    assert Path(docx_path).stat().st_size > 0


def test_materialize_original_docx_from_uploaded_resume(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))
    monkeypatch.setattr(Config, "DB_PATH", str(tmp_path / "career_path.db"))

    import base64
    import io
    import docx

    db = Database(str(tmp_path / "auth.db"))
    user_id = "user-docx"
    resume_id = "resume-docx"
    task_id = "materialize-docx-task"

    document = docx.Document()
    document.add_paragraph("Original template paragraph")
    buffer = io.BytesIO()
    document.save(buffer)
    docx_bytes = buffer.getvalue()

    db.save_user_resume(
        user_id=user_id,
        resume_id=resume_id,
        file_name="template_resume.docx",
        file_size=len(docx_bytes),
        file_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        parsed_resume={
            "cleanedText": "Original template paragraph",
            "file": {
                "name": "template_resume.docx",
                "b64_content": base64.b64encode(docx_bytes).decode("ascii"),
            },
        },
    )

    result = materialize_original_resume_file(user_id, task_id, resume_id, db=db)

    original_docx = Path(get_safe_workspace_path(user_id, task_id, "original_resume.docx"))
    assert result == {"source_format": "docx", "path": str(original_docx)}
    assert original_docx.exists()
    assert original_docx.read_bytes() == docx_bytes


def test_materialize_original_pdf_from_uploaded_resume(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))
    monkeypatch.setattr(Config, "DB_PATH", str(tmp_path / "career_path.db"))

    import base64

    db = Database(str(tmp_path / "auth.db"))
    user_id = "user-pdf"
    resume_id = "resume-pdf"
    task_id = "materialize-pdf-task"
    pdf_bytes = b"%PDF-1.4\n% internpath fixture\n"

    db.save_user_resume(
        user_id=user_id,
        resume_id=resume_id,
        file_name="template_resume.pdf",
        file_size=len(pdf_bytes),
        file_type="application/pdf",
        parsed_resume={
            "cleanedText": "PDF resume text",
            "file": {
                "name": "template_resume.pdf",
                "b64_content": base64.b64encode(pdf_bytes).decode("ascii"),
            },
        },
    )

    result = materialize_original_resume_file(user_id, task_id, resume_id, db=db)

    original_pdf = Path(get_safe_workspace_path(user_id, task_id, "original_resume.pdf"))
    assert result == {"source_format": "pdf", "path": str(original_pdf)}
    assert original_pdf.exists()
    assert original_pdf.read_bytes() == pdf_bytes


def test_docx_template_fingerprint_detects_lost_media_and_relationships(tmp_path):
    import re
    import zipfile

    import docx
    from PIL import Image

    original_docx = tmp_path / "original.docx"
    broken_docx = tmp_path / "broken.docx"
    image_path = tmp_path / "photo.png"

    Image.new("RGB", (12, 12), "blue").save(image_path)

    document = docx.Document()
    document.sections[0].header.add_paragraph("Header text")
    document.sections[0].footer.add_paragraph("Footer text")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Table cell"
    document.add_paragraph("Body text")
    document.add_picture(str(image_path))
    document.save(original_docx)

    with zipfile.ZipFile(original_docx, "r") as source, zipfile.ZipFile(broken_docx, "w") as target:
        for info in source.infolist():
            if info.filename.startswith("word/media/"):
                continue
            data = source.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                text = data.decode("utf-8")
                data = re.sub(r"<Relationship[^>]+/image[^>]*/>", "", text).encode("utf-8")
            target.writestr(info, data)

    before = collect_docx_template_fingerprint(str(original_docx))
    after = collect_docx_template_fingerprint(str(broken_docx))
    report = validate_docx_template_preservation(before, after)

    assert before["media_files"] > after["media_files"]
    assert before["image_relationships"] > after["image_relationships"]
    assert report["ok"] is False
    assert any(issue["field"] == "media_files" for issue in report["issues"])
    assert any(issue["field"] == "image_relationships" for issue in report["issues"])


def test_update_docx_resume_from_log_removes_unmatched_output(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))

    import docx

    user_id = "user-docx-no-match"
    task_id = "no-match-docx"
    original_docx = Path(get_safe_workspace_path(user_id, task_id, "original_resume.docx"))
    optimized_docx = Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx"))
    mod_log_path = Path(get_safe_workspace_path(user_id, task_id, "modification_log.json"))

    document = docx.Document()
    document.add_paragraph("Only original body text")
    document.save(original_docx)

    mod_log_path.write_text(
        json.dumps(
            [{
                "section_name": "missing",
                "section_index": 1,
                "original": "This text is not in the original docx",
                "new": "Replacement text",
                "reason": "test no lossy success",
            }],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    assert update_docx_resume_from_log(user_id, task_id) is False
    assert not optimized_docx.exists()


def test_update_docx_resume_from_log_rejects_template_structure_loss(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))

    import docx
    import backend.agents.tools.docx_tools as docx_tools

    user_id = "user-docx-guard-loss"
    task_id = "guard-loss-docx"
    original_docx = Path(get_safe_workspace_path(user_id, task_id, "original_resume.docx"))
    optimized_docx = Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx"))
    guard_report = Path(get_safe_workspace_path(user_id, task_id, "docx_template_guard.json"))
    mod_log_path = Path(get_safe_workspace_path(user_id, task_id, "modification_log.json"))

    document = docx.Document()
    document.add_paragraph("Original body text")
    document.save(original_docx)

    mod_log_path.write_text(
        json.dumps(
            [{
                "section_name": "body",
                "section_index": 1,
                "original": "Original body text",
                "new": "Updated body text",
                "reason": "test template guard",
            }],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    real_collect = docx_tools.collect_docx_template_fingerprint

    def fake_collect(path: str) -> dict[str, int]:
        fingerprint = real_collect(path)
        if str(path).endswith("optimized_resume.docx"):
            fingerprint["paragraphs"] = 0
        return fingerprint

    monkeypatch.setattr(docx_tools, "collect_docx_template_fingerprint", fake_collect)

    assert update_docx_resume_from_log(user_id, task_id) is False
    assert not optimized_docx.exists()
    report = json.loads(guard_report.read_text(encoding="utf-8"))
    assert report["ok"] is False
    assert any(issue["field"] == "paragraphs" for issue in report["issues"])


def test_update_docx_resume_from_log_matches_split_docx_paragraphs(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))

    import docx

    user_id = "user-docx-replace"
    task_id = "split-paragraph-docx"
    original_docx = Path(get_safe_workspace_path(user_id, task_id, "original_resume.docx"))
    optimized_docx = Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx"))
    mod_log_path = Path(get_safe_workspace_path(user_id, task_id, "modification_log.json"))

    document = docx.Document()
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "模板侧栏"
    document.add_paragraph("项目经历")
    document.add_paragraph("InternPath 简历优化工作台")
    document.add_paragraph("负责开发 Agent 流程。")
    document.save(original_docx)

    original_text = "项目经历\nInternPath 简历优化工作台\n负责开发 Agent 流程。"
    new_text = "项目经历\nInternPath 简历优化工作台\n负责开发流式 Agent 简历优化链路。"
    mod_log_path.write_text(
        json.dumps(
            [{
                "section_name": "项目经历",
                "section_index": 1,
                "original": original_text,
                "new": new_text,
                "reason": "匹配 JD",
            }],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    assert update_docx_resume_from_log(user_id, task_id) is True
    assert optimized_docx.exists()

    updated = docx.Document(optimized_docx)
    body_text = "\n".join(paragraph.text for paragraph in updated.paragraphs)
    assert "负责开发流式 Agent 简历优化链路。" in body_text
    assert "负责开发 Agent 流程。" not in body_text
    assert updated.tables[0].cell(0, 0).text == "模板侧栏"


def test_update_docx_resume_from_log_inserts_extra_paragraphs(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))

    import docx

    user_id = "user-docx-extra-lines"
    task_id = "extra-line-docx"
    original_docx = Path(get_safe_workspace_path(user_id, task_id, "original_resume.docx"))
    optimized_docx = Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx"))
    mod_log_path = Path(get_safe_workspace_path(user_id, task_id, "modification_log.json"))

    document = docx.Document()
    document.add_paragraph("负责开发 Agent 流程。")
    document.save(original_docx)

    mod_log_path.write_text(
        json.dumps(
            [{
                "section_name": "项目经历",
                "section_index": 1,
                "original": "负责开发 Agent 流程。",
                "new": "负责开发流式 Agent 简历优化链路。\n支撑 Responses API 缓存诊断。\n沉淀 DOCX 高保真替换能力。",
                "reason": "匹配 JD",
            }],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    assert update_docx_resume_from_log(user_id, task_id) is True
    updated = docx.Document(optimized_docx)
    paragraph_texts = [paragraph.text for paragraph in updated.paragraphs]

    assert "负责开发流式 Agent 简历优化链路。" in paragraph_texts
    assert "支撑 Responses API 缓存诊断。" in paragraph_texts
    assert "沉淀 DOCX 高保真替换能力。" in paragraph_texts
    assert all("\n" not in text for text in paragraph_texts)


def test_update_docx_resume_from_log_matches_vml_textbox_paragraphs(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))

    import docx
    import zipfile
    from docx.oxml import parse_xml

    user_id = "user-docx-textbox-replace"
    task_id = "textbox-docx"
    original_docx = Path(get_safe_workspace_path(user_id, task_id, "original_resume.docx"))
    optimized_docx = Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx"))
    mod_log_path = Path(get_safe_workspace_path(user_id, task_id, "modification_log.json"))

    document = docx.Document()
    textbox_paragraph = parse_xml(
        """
        <w:p
          xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
          xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
          xmlns:v="urn:schemas-microsoft-com:vml">
          <w:r>
            <w:pict>
              <v:shape style="height:50pt">
                <v:textbox>
                  <w:txbxContent>
                    <w:p><w:r><w:t>Textbox original content</w:t></w:r></w:p>
                  </w:txbxContent>
                </v:textbox>
              </v:shape>
            </w:pict>
          </w:r>
        </w:p>
        """
    )
    document._body._element.append(textbox_paragraph)
    document.save(original_docx)

    mod_log_path.write_text(
        json.dumps(
            [{
                "section_name": "textbox",
                "section_index": 1,
                "original": "Textbox original content",
                "new": "Textbox updated content",
                "reason": "keep template textbox",
            }],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    assert update_docx_resume_from_log(user_id, task_id) is True
    assert optimized_docx.exists()

    with zipfile.ZipFile(optimized_docx) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "Textbox updated content" in document_xml
    assert "Textbox original content" not in document_xml


def test_update_docx_resume_from_log_matches_long_textbox_windows(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))

    import docx
    import zipfile
    from docx.oxml import parse_xml

    user_id = "user-docx-long-textbox-replace"
    task_id = "long-textbox-docx"
    original_docx = Path(get_safe_workspace_path(user_id, task_id, "original_resume.docx"))
    optimized_docx = Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx"))
    mod_log_path = Path(get_safe_workspace_path(user_id, task_id, "modification_log.json"))
    parts = [f"part-{idx:02d}" for idx in range(45)]

    document = docx.Document()
    for part in parts:
        textbox_paragraph = parse_xml(
            f"""
            <w:p
              xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
              xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
              xmlns:v="urn:schemas-microsoft-com:vml">
              <w:r>
                <w:pict>
                  <v:shape style="height:12pt">
                    <v:textbox>
                      <w:txbxContent>
                        <w:p><w:r><w:t>{part}</w:t></w:r></w:p>
                      </w:txbxContent>
                    </v:textbox>
                  </v:shape>
                </w:pict>
              </w:r>
            </w:p>
            """
        )
        document._body._element.append(textbox_paragraph)
    document.save(original_docx)

    mod_log_path.write_text(
        json.dumps(
            [{
                "section_name": "long textbox",
                "section_index": 1,
                "original": "\n".join(parts),
                "new": "long textbox updated content",
                "reason": "keep multi-box template",
            }],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    assert update_docx_resume_from_log(user_id, task_id) is True
    assert optimized_docx.exists()

    with zipfile.ZipFile(optimized_docx) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "long textbox updated content" in document_xml
    assert "part-44" not in document_xml


def test_api_download_rejects_lossy_docx_fallback(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "EMAIL_VERIFICATION_REQUIRED", False)
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path / "user_data"))
    monkeypatch.setattr(Config, "DB_PATH", str(tmp_path / "career_path.db"))

    # Setup service/app
    service = object.__new__(CareerPathAIService)
    db = Database(str(tmp_path / "auth.db"))
    app = create_app(service=service, auth_db=db)
    client = TestClient(app)

    # Register and create task
    reg = client.post("/api/auth/register", json={"username": "demo@example.com", "password": "password123"})
    token = reg.json()["token"]
    user_id = reg.json()["user"]["id"]
    task_id = "test_api_task"

    # Create the task in DB first
    db.create_agent_resume_task(
        task_id=task_id,
        user_id=user_id,
        resume_id="fake_resume_id",
        original_resume_name="my_resume.md",
        jd_text="fake jd text"
    )

    # Write a fake completed agent task to DB
    db.update_agent_resume_task_status(
        task_id=task_id,
        user_id=user_id,
        status="COMPLETED",
        logs="[]",
        optimized_resume_md="### Optimized Resume"
    )
    # Also update original_resume_name
    conn = db.get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE agent_resume_tasks SET original_resume_name = ? WHERE task_id = ?",
        ("my_resume.md", task_id)
    )
    conn.commit()

    # Test MD download
    resp = client.get(f"/api/agent/resume/tasks/{task_id}/download", headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "text/markdown; charset=utf-8"
    assert "optimized_my_resume.md" in resp.headers["content-disposition"]

    # DOCX export must not fabricate a plain-text DOCX when no high-fidelity output exists.
    resp_docx_missing = client.get(
        f"/api/agent/resume/tasks/{task_id}/download?format=docx",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert resp_docx_missing.status_code == 409
    assert "高保真" in resp_docx_missing.json()["detail"]
    assert not Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx")).exists()

    resp_pdf_missing = client.get(
        f"/api/agent/resume/tasks/{task_id}/download?format=pdf",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert resp_pdf_missing.status_code == 409
    assert "高保真" in resp_pdf_missing.json()["detail"]
    assert not Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx")).exists()

    # Now simulate docx creation in workspace
    workspace_dir = Path(get_safe_workspace_path(user_id, task_id, "optimized_resume.docx")).parent
    with open(workspace_dir / "optimized_resume.docx", "w") as f:
        f.write("fake docx content")
    with open(workspace_dir / "modification_log.json", "w", encoding="utf-8") as f:
        json.dump([{"section_name": "项目经历", "original": "旧内容", "new": "新内容", "reason": "匹配 JD"}], f, ensure_ascii=False)

    resp_docx_success = client.get(
        f"/api/agent/resume/tasks/{task_id}/download?format=docx",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert resp_docx_success.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in resp_docx_success.headers["content-type"]
    assert "optimized_my_resume.docx" in resp_docx_success.headers["content-disposition"]

    resp_list = client.get("/api/agent/resume/tasks", headers={"Authorization": f"Bearer {token}"})
    assert resp_list.status_code == 200
    listed_task = resp_list.json()["tasks"][0]
    assert listed_task["hasDocx"] is True
    assert listed_task["modificationLog"][0]["section_name"] == "项目经历"
