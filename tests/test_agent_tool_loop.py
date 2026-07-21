import json
from unittest.mock import MagicMock

from backend.agents.tool_loop import AgenticToolLoop, parse_json_action
from backend.agents.tool_loop import ResponsesJsonActionModelTurn
from backend.agents.tool_loop import ResponsesNativeToolModelTurn, ResponsesNativeTurnResult
from backend.agents.schemas import extract_json_object
from backend.agents.tool_registry import AgentToolContext, AgentToolRegistry
from backend.agents.tools.hitl_tool import HumanInteractionRequired
from backend.agents.tools.workspace_tools import tool_read_file, tool_write_file
from config import Config


def test_parse_json_action_accepts_fenced_json():
    action = parse_json_action(
        """```json
        {"action":"call_tool","tool":"list_workspace_files","arguments":{}}
        ```"""
    )

    assert action["action"] == "call_tool"
    assert action["tool"] == "list_workspace_files"


def test_parse_json_action_accepts_trailing_model_text():
    action = parse_json_action(
        """{"action":"call_tool","tool":"list_workspace_files","arguments":{}}

I will inspect the workspace before editing."""
    )

    assert action["action"] == "call_tool"
    assert action["tool"] == "list_workspace_files"


def test_parse_json_action_accepts_prefixed_model_text():
    action = parse_json_action(
        """Here is the next action:
{"action":"final_answer","content":"done"}"""
    )

    assert action["action"] == "final_answer"
    assert action["content"] == "done"


def test_extract_json_object_accepts_trailing_model_text():
    payload = extract_json_object(
        """{"steps":[{"step_index":1,"section_index":0,"section_name":"Projects","original_content":"Old","improvement_goal":"Improve","status":"PENDING"}]}

Plan summary: next I will rewrite the Projects section."""
    )

    assert payload["steps"][0]["section_name"] == "Projects"


def test_json_action_loop_calls_tools_and_finishes(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    monkeypatch.setattr("backend.agents.tool_registry.convert_docx_to_pdf", lambda *_args, **_kwargs: False)
    import docx

    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=5)
    ctx = AgentToolContext(
        user_id="u1",
        task_id="t1",
        resume_text="Work Experience\nBuilt APIs",
        jd_text="Backend engineer",
    )
    doc = docx.Document()
    doc.add_paragraph("Work Experience")
    doc.add_paragraph("Built APIs")
    original_docx_path = tmp_path / "workspaces" / "user_u1" / "task_t1" / "original_resume.docx"
    original_docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(original_docx_path))
    replies = iter([
        '{"action":"call_tool","tool":"extract_resume_sections","arguments":{}}',
        '{"action":"call_tool","tool":"replace_resume_section","arguments":{"section_index":0,"new_content":"Work Experience\\nBuilt reliable backend APIs","reason":"backend JD"}}',
        '{"action":"call_tool","tool":"generate_modification_diff","arguments":{}}',
        '{"action":"call_tool","tool":"finalize_resume_artifacts","arguments":{}}',
        '{"action":"final_answer","content":"done"}',
    ])

    result = loop.run_json_action_loop(ctx, lambda _messages: next(replies))

    assert result.status == "completed"
    assert result.final_text == "done"
    assert any(event["type"] == "tool_call" and event["tool_name"] == "replace_resume_section" for event in result.events)
    assembled = tool_read_file(ctx.user_id, ctx.task_id, "assembled_resume.txt")
    assert "Built reliable backend APIs" in assembled


def test_json_action_loop_bootstraps_workspace_tools_before_model_turn(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=1)
    ctx = AgentToolContext(
        user_id="u1",
        task_id="t-bootstrap",
        resume_text="Projects\nBuilt APIs",
        jd_text="Backend engineer",
    )
    tool_write_file(ctx.user_id, ctx.task_id, "original_resume.txt", ctx.resume_text)
    tool_write_file(ctx.user_id, ctx.task_id, "job_description.txt", ctx.jd_text)
    messages_seen = []

    def model_turn(messages):
        messages_seen.append([dict(item) for item in messages])
        return '{"action":"final_answer","content":"done"}'

    result = loop.run_json_action_loop(ctx, model_turn)

    assert result.status == "max_turns"
    first_model_event_index = next(index for index, event in enumerate(result.events) if event["type"] == "model_turn_start")
    bootstrap_events = result.events[:first_model_event_index]
    assert any(event["type"] == "tool_call" and event["tool_name"] == "list_workspace_files" for event in bootstrap_events)
    assert any(
        event["type"] == "tool_call"
        and event["tool_name"] == "read_workspace_file"
        and event["arguments"] == {"filename": "original_resume.txt"}
        for event in bootstrap_events
    )
    assert any(event["type"] == "tool_call" and event["tool_name"] == "extract_resume_sections" for event in bootstrap_events)
    assert all(event.get("bootstrap") is True for event in bootstrap_events if event["type"] in {"tool_call", "tool_result"})
    assert "bootstrap_tool_results" in messages_seen[0][-1]["content"]
    assert "Projects" in tool_read_file(ctx.user_id, ctx.task_id, "resume_sections.json")
    assert any(
        event["type"] == "error"
        and "replace_resume_section" in event["error"]
        for event in result.events
    )


def test_json_action_loop_feedbacks_invalid_json_and_recovers(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=3)
    ctx = AgentToolContext(user_id="u1", task_id="t1")
    messages_seen = []
    replies = iter([
        "not json",
        '{"action":"final_answer","content":"fixed"}',
    ])

    def model_turn(messages):
        messages_seen.append(messages[-1]["content"])
        return next(replies)

    result = loop.run_json_action_loop(ctx, model_turn)

    assert result.status == "completed"
    assert result.final_text == "fixed"
    assert any(event["type"] == "error" and event["retryable"] is True for event in result.events)
    assert "not a valid JSON action" in messages_seen[-1]


def test_json_action_loop_stops_at_max_turns(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=2)
    ctx = AgentToolContext(user_id="u1", task_id="t1")

    result = loop.run_json_action_loop(
        ctx,
        lambda _messages: '{"action":"call_tool","tool":"list_workspace_files","arguments":{}}',
    )

    assert result.status == "max_turns"
    assert result.error == "max_turns exceeded"


def test_json_action_loop_returns_needs_human_instead_of_raising(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    question = "请补充项目的量化结果。"

    class _Registry:
        def list_tools(self, *, is_co_pilot=True):
            return []

        def execute(self, tool_name, arguments, _ctx):
            if tool_name == "ask_user_for_fact":
                raise HumanInteractionRequired(arguments["question"])
            return {
                "tool_name": tool_name,
                "ok": True,
                "result": {},
                "error": "",
                "duration_ms": 0,
            }

    loop = AgenticToolLoop(_Registry(), max_turns=1)
    ctx = AgentToolContext(user_id="u1", task_id="t-hitl")

    result = loop.run_json_action_loop(
        ctx,
        lambda _messages: json.dumps({
            "action": "call_tool",
            "tool": "ask_user_for_fact",
            "arguments": {"question": question},
        }, ensure_ascii=False),
    )

    assert result.status == "needs_human"
    assert result.final_text == question
    assert any(
        event.get("type") == "tool_call"
        and event.get("tool_name") == "ask_user_for_fact"
        for event in result.events
    )


def test_json_action_loop_replays_failed_tool_before_model_turn(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=1)
    ctx = AgentToolContext(user_id="u1", task_id="t1")
    failure_point = {
        "failed_tool_name": "write_workspace_file",
        "failed_tool_arguments": {"filename": "retry.txt", "content": "same payload"},
        "failed_sequence": 7,
        "retry_count": 2,
    }
    messages_seen = []

    def model_turn(messages):
        messages_seen.append(messages)
        return '{"action":"final_answer","content":"continued"}'

    result = loop.run_json_action_loop(ctx, model_turn, retry_failure_point=failure_point)

    assert result.status == "completed"
    assert any(
        event["type"] == "tool_call"
        and event["tool_name"] == "write_workspace_file"
        and event["arguments"] == failure_point["failed_tool_arguments"]
        and event["retry"] is True
        for event in result.events
    )
    assert any(
        event["type"] == "tool_result"
        and event["tool_name"] == "write_workspace_file"
        and event["retry"] is True
        for event in result.events
    )
    replay_message = messages_seen[0][-1]["content"]
    assert "retry_tool_result" in replay_message
    assert "same payload" in replay_message


def test_json_action_loop_records_and_reuses_failed_model_input(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=3)
    ctx = AgentToolContext(user_id="u1", task_id="t1")
    replies = iter([
        '{"action":"call_tool","tool":"list_workspace_files","arguments":{}}',
    ])

    def flaky_model_turn(messages):
        try:
            return next(replies)
        except StopIteration as exc:
            raise RuntimeError("provider disconnected") from exc

    failed = loop.run_json_action_loop(ctx, flaky_model_turn)

    assert failed.status == "failed"
    error_event = next(event for event in failed.events if event["type"] == "error")
    failed_model_input = error_event["failed_model_input"]
    assert failed_model_input["mode"] == "json_action"
    assert failed_model_input["messages"]
    assert error_event["last_successful_tool_result"]["tool_name"] == "list_workspace_files"

    retry_messages_seen = []

    def retry_model_turn(messages):
        retry_messages_seen.append(messages)
        return '{"action":"final_answer","content":"continued"}'

    retried = loop.run_json_action_loop(
        ctx,
        retry_model_turn,
        retry_failure_point={
            "failed_stage": "model_turn",
            "failed_model_input": failed_model_input,
        },
    )

    assert retried.status == "completed"
    assert retry_messages_seen[0] == failed_model_input["messages"]


def test_responses_json_action_model_turn_uses_stream_and_records_usage(monkeypatch):
    monkeypatch.setattr(Config, "OPENAI_PROMPT_CACHE_STABLE_PREFIX_ENABLED", True)
    client = MagicMock()
    client._internpath_prompt_cache_key = "agentic-cache-key"
    client._internpath_prompt_cache_retention = "24h"
    client.responses.create.return_value = [
        {"type": "response.output_text.delta", "delta": '{"action":"final_answer",'},
        {"type": "response.output_text.delta", "delta": '"content":"done"}'},
        {
            "type": "response.completed",
            "response": {
                "usage": {
                    "input_tokens": 1300,
                    "output_tokens": 12,
                    "input_tokens_details": {"cached_tokens": 512},
                }
            },
        },
    ]
    stream_deltas = []
    model_turn = ResponsesJsonActionModelTurn(
        client,
        "gpt-test",
        namespace="agentic_resume_json_action",
        stream_delta_sink=stream_deltas.append,
    )

    text = model_turn([
        {"role": "system", "content": "Return JSON actions only."},
        {"role": "user", "content": "finish"},
    ])

    assert text == '{"action":"final_answer","content":"done"}'
    kwargs = client.responses.create.call_args.kwargs
    assert kwargs["model"] == "gpt-test"
    assert kwargs["stream"] is True
    assert kwargs["prompt_cache_key"] == "agentic-cache-key"
    assert kwargs["prompt_cache_retention"] == "24h"
    assert "Return JSON actions only." in kwargs["instructions"]
    assert "InternPath" in kwargs["instructions"]
    assert model_turn.deltas == ['{"action":"final_answer",', '"content":"done"}']
    assert stream_deltas == model_turn.deltas
    assert model_turn.last_provider_usage["providerEndpointMode"] == "responses"
    assert model_turn.last_provider_usage["providerStream"] is True
    assert model_turn.last_provider_usage["providerCachedTokens"] == 512
    client.chat.completions.create.assert_not_called()


def test_json_action_loop_emits_provider_usage_from_model_turn(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    client = MagicMock()
    client._internpath_prompt_cache_key = "agentic-cache-key"
    client._internpath_prompt_cache_retention = "24h"
    client.responses.create.return_value = [
        {"type": "response.output_text.delta", "delta": '{"action":"final_answer","content":"done"}'},
        {
            "type": "response.completed",
            "response": {
                "usage": {
                    "input_tokens": 1300,
                    "output_tokens": 12,
                    "input_tokens_details": {"cached_tokens": 256},
                }
            },
        },
    ]
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=1)
    ctx = AgentToolContext(user_id="u1", task_id="t1")
    model_turn = ResponsesJsonActionModelTurn(client, "gpt-test")

    result = loop.run_json_action_loop(ctx, model_turn)

    assert result.status == "completed"
    assert any(event["type"] == "model_stream_delta" and event["content"] for event in result.events)
    assert any(
        event["type"] == "provider_usage" and event["providerCachedTokens"] == 256
        for event in result.events
    )


def test_native_responses_loop_executes_function_call_and_continues(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    monkeypatch.setattr(Config, "OPENAI_PROMPT_CACHE_STABLE_PREFIX_ENABLED", True)
    monkeypatch.setattr("backend.agents.tool_registry.convert_docx_to_pdf", lambda *_args, **_kwargs: False)
    import docx

    client = MagicMock()
    client._internpath_prompt_cache_key = "agentic-cache-key"
    client._internpath_prompt_cache_retention = "24h"
    client.responses.create.side_effect = [
        [
            {
                "type": "response.completed",
                "response": {
                    "id": "resp_tool",
                    "output": [
                        {
                            "type": "function_call",
                            "name": "replace_resume_section",
                            "arguments": json.dumps({
                                "section_index": 0,
                                "new_content": "Projects\nBuilt reliable backend APIs",
                                "reason": "Backend JD",
                            }),
                            "call_id": "call_1",
                        }
                    ],
                    "usage": {
                        "input_tokens": 1400,
                        "output_tokens": 16,
                        "input_tokens_details": {"cached_tokens": 300},
                    },
                },
            }
        ],
        [
            {
                "type": "response.completed",
                "response": {
                    "id": "resp_finalize",
                    "output": [
                        {
                            "type": "function_call",
                            "name": "finalize_resume_artifacts",
                            "arguments": "{}",
                            "call_id": "call_2",
                        }
                    ],
                    "usage": {
                        "input_tokens": 240,
                        "output_tokens": 8,
                        "input_tokens_details": {"cached_tokens": 180},
                    },
                },
            }
        ],
        [
            {"type": "response.output_text.delta", "delta": "done"},
            {
                "type": "response.completed",
                "response": {
                    "id": "resp_done",
                    "output": [
                        {
                            "type": "message",
                            "role": "assistant",
                            "content": [{"type": "output_text", "text": "done"}],
                        }
                    ],
                    "usage": {
                        "input_tokens": 120,
                        "output_tokens": 4,
                        "input_tokens_details": {"cached_tokens": 120},
                    },
                },
            },
        ],
    ]
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=3)
    ctx = AgentToolContext(user_id="u-native", task_id="t-native", resume_text="Projects\nBuilt APIs", jd_text="Backend")
    original_docx_path = tmp_path / "workspaces" / "user_u-native" / "task_t-native" / "original_resume.docx"
    original_docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()
    doc.add_paragraph("Projects")
    doc.add_paragraph("Built APIs")
    doc.save(str(original_docx_path))
    model_turn = ResponsesNativeToolModelTurn(
        client,
        "gpt-test",
        tools=loop.registry.responses_tools(is_co_pilot=True),
    )

    result = loop.run_native_responses_loop(ctx, model_turn)

    assert result.status == "completed"
    assert result.final_text == "done"
    assert client.responses.create.call_count == 3
    first_kwargs = client.responses.create.call_args_list[0].kwargs
    assert first_kwargs["model"] == "gpt-test"
    assert first_kwargs["stream"] is True
    assert first_kwargs["tools"]
    assert first_kwargs["prompt_cache_key"] == "agentic-cache-key"
    assert first_kwargs["prompt_cache_retention"] == "24h"
    second_kwargs = client.responses.create.call_args_list[1].kwargs
    assert second_kwargs["previous_response_id"] == "resp_tool"
    assert second_kwargs["input"][0]["type"] == "function_call_output"
    assert second_kwargs["input"][0]["call_id"] == "call_1"
    tool_output = json.loads(second_kwargs["input"][0]["output"])
    assert tool_output["tool_name"] == "replace_resume_section"
    assert tool_output["ok"] is True
    third_kwargs = client.responses.create.call_args_list[2].kwargs
    assert third_kwargs["previous_response_id"] == "resp_finalize"
    finalize_output = json.loads(third_kwargs["input"][0]["output"])
    assert finalize_output["tool_name"] == "finalize_resume_artifacts"
    assert finalize_output["ok"] is True
    assert any(event["type"] == "tool_call" and event["tool_name"] == "replace_resume_section" for event in result.events)
    assert any(event["type"] == "tool_call" and event["tool_name"] == "finalize_resume_artifacts" for event in result.events)
    assert any(event["type"] == "model_stream_delta" and event["content"] == "done" for event in result.events)
    assert any(event["type"] == "provider_usage" and event["providerEndpointMode"] == "responses" for event in result.events)
    client.chat.completions.create.assert_not_called()


def test_native_responses_loop_records_and_reuses_failed_model_input(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=2)
    ctx = AgentToolContext(user_id="u-native-retry", task_id="t-native-retry")

    class FailingNativeTurn:
        last_provider_usage = {}

        def create(self, input_items, *, instructions="", previous_response_id="", on_delta=None):
            raise RuntimeError("responses stream disconnected")

    failed = loop.run_native_responses_loop(ctx, FailingNativeTurn())

    assert failed.status == "failed"
    error_event = next(event for event in failed.events if event["type"] == "error")
    failed_model_input = error_event["failed_model_input"]
    assert failed_model_input["mode"] == "native_responses"
    assert failed_model_input["input_items"]
    assert "instructions" in failed_model_input

    class RetryNativeTurn:
        last_provider_usage = {}

        def __init__(self):
            self.calls = []

        def create(self, input_items, *, instructions="", previous_response_id="", on_delta=None):
            self.calls.append({
                "input_items": input_items,
                "instructions": instructions,
                "previous_response_id": previous_response_id,
            })
            return ResponsesNativeTurnResult(text="continued", response={"id": "resp_done", "output": []}, response_id="resp_done")

    retry_turn = RetryNativeTurn()
    retried = loop.run_native_responses_loop(
        ctx,
        retry_turn,
        retry_failure_point={
            "failed_stage": "model_turn",
            "failed_model_input": failed_model_input,
        },
    )

    assert retried.status == "completed"
    assert retry_turn.calls[0]["input_items"] == failed_model_input["input_items"]
    assert retry_turn.calls[0]["instructions"] == failed_model_input["instructions"]
    assert retry_turn.calls[0]["previous_response_id"] == failed_model_input["previous_response_id"]


def test_native_responses_loop_returns_structured_invalid_argument_feedback(tmp_path, monkeypatch):
    monkeypatch.setattr(Config, "USER_DB_DIR", str(tmp_path))
    loop = AgenticToolLoop(AgentToolRegistry(), max_turns=1)
    ctx = AgentToolContext(user_id="u-invalid-native", task_id="t-invalid-native")

    class InvalidArgumentsNativeTurn:
        last_provider_usage = {}

        def create(self, _input_items, *, instructions="", previous_response_id="", on_delta=None):
            return ResponsesNativeTurnResult(
                response={
                    "id": "resp-invalid",
                    "output": [{
                        "type": "function_call",
                        "name": "list_workspace_files",
                        "arguments": "[]",
                        "call_id": "call-invalid",
                    }],
                },
                response_id="resp-invalid",
            )

    result = loop.run_native_responses_loop(ctx, InvalidArgumentsNativeTurn())

    assert result.status == "max_turns"
    failure = next(
        event
        for event in result.events
        if event["type"] == "tool_result" and event.get("failure_code") == "invalid_tool_arguments"
    )
    assert failure["ok"] is False
    assert failure["failure_code"] == "invalid_tool_arguments"
    assert failure["failure_reason"]
    assert failure["retryable"] is True
